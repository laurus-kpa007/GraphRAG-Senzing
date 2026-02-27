"""
Integration tests for GraphRAG-Senzing pipeline.
Run: python -m pytest tests/ -v
Or:  python tests/test_integration.py
"""

import json
import pathlib
import shutil
import sys
import unittest

import numpy as np

# Add project root to path
sys.path.insert(0, str(pathlib.Path(__file__).parent.parent))

from src.loaders import DocumentLoader
from src.embeddings import OllamaEmbedding, OllamaLLM
from src.pipeline import AgenticPipeline


class TestDocumentLoader(unittest.TestCase):
    def setUp(self):
        self.loader = DocumentLoader()
        self.docs_dir = pathlib.Path("data/documents")

    def test_load_markdown(self):
        paras = self.loader.load(self.docs_dir / "sample.md")
        self.assertGreater(len(paras), 0)
        self.assertIsInstance(paras[0], str)

    def test_load_text(self):
        paras = self.loader.load(self.docs_dir / "sample.txt")
        self.assertGreater(len(paras), 0)

    def test_load_korean_markdown(self):
        paras = self.loader.load(self.docs_dir / "sample_ko.md")
        self.assertGreater(len(paras), 0)
        # Verify Korean chars preserved
        has_korean = any("\uAC00" <= c <= "\uD7A3" for p in paras for c in p)
        self.assertTrue(has_korean, "Korean characters should be preserved")

    def test_load_korean_text(self):
        paras = self.loader.load(self.docs_dir / "sample_ko.txt")
        self.assertGreater(len(paras), 0)

    def test_load_directory(self):
        all_docs = self.loader.load_directory(self.docs_dir)
        self.assertGreaterEqual(len(all_docs), 4)

    def test_load_docx(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Test paragraph")
        doc.add_paragraph("한글 테스트")

        test_path = self.docs_dir / "_test.docx"
        doc.save(str(test_path))
        try:
            paras = self.loader.load(test_path)
            self.assertEqual(len(paras), 2)
            self.assertIn("한글", paras[1])
        finally:
            test_path.unlink(missing_ok=True)

    def test_unsupported_format(self):
        with self.assertRaises(ValueError):
            self.loader.load("config.toml")

    def test_file_not_found(self):
        with self.assertRaises(FileNotFoundError):
            self.loader.load("nonexistent.txt")

    def test_scrub_korean_nfc(self):
        # NFC should preserve composed Korean
        result = DocumentLoader._scrub("안녕하세요")
        self.assertEqual(result, "안녕하세요")


class TestOllamaClients(unittest.TestCase):
    def test_embedding_init(self):
        embed = OllamaEmbedding(model="bona/bge-m3:latest", dim=1024)
        self.assertEqual(embed.model, "bona/bge-m3:latest")
        self.assertEqual(embed.dim, 1024)

    def test_llm_init(self):
        llm = OllamaLLM(model="gemma3:27b")
        self.assertEqual(llm.model, "gemma3:27b")
        self.assertEqual(llm.temperature, 0.0)

    def test_embed_availability_no_server(self):
        embed = OllamaEmbedding(base_url="http://localhost:99999")
        self.assertFalse(embed.is_available())

    def test_llm_availability_no_server(self):
        llm = OllamaLLM(base_url="http://localhost:99999")
        self.assertFalse(llm.is_available())

    def test_embed_error_handling(self):
        embed = OllamaEmbedding(base_url="http://localhost:99999")
        with self.assertRaises(Exception):
            embed.embed_text("test")

    def test_llm_error_handling(self):
        llm = OllamaLLM(base_url="http://localhost:99999")
        with self.assertRaises(Exception):
            llm.generate("test")


class TestPipeline(unittest.TestCase):
    def setUp(self):
        self.pipeline = AgenticPipeline()
        # Clean test data
        for d in ["data/lancedb", "data/output", "data/cache"]:
            p = pathlib.Path(d)
            if p.exists():
                shutil.rmtree(p)

    def tearDown(self):
        for d in ["data/lancedb", "data/output", "data/cache"]:
            p = pathlib.Path(d)
            if p.exists():
                shutil.rmtree(p)

    def test_init(self):
        self.assertIsNotNone(self.pipeline.config)
        self.assertIn("rag", self.pipeline.config)
        self.assertIn("embed", self.pipeline.config)

    def test_check_prerequisites(self):
        status = self.pipeline.check_prerequisites()
        self.assertIn("ollama_server", status)
        self.assertIn("llm_model", status)
        self.assertIn("embed_model", status)

    def test_initialize(self):
        self.pipeline.initialize()
        self.assertTrue(self.pipeline._is_initialized)
        self.assertIsNotNone(self.pipeline._lance_db)

    def test_load_documents(self):
        docs = self.pipeline.load_documents(["data/documents/"])
        self.assertGreaterEqual(len(docs), 4)

    def test_make_chunks(self):
        paras = ["Short paragraph."] * 20
        chunks = self.pipeline.make_chunks(paras, max_chunk_size=100)
        self.assertGreater(len(chunks), 1)
        for chunk in chunks:
            self.assertLessEqual(len(chunk), 200)  # Some margin

    def test_language_detection(self):
        self.assertEqual(self.pipeline._detect_language("Hello world"), "en")
        self.assertEqual(self.pipeline._detect_language("안녕하세요"), "ko")
        self.assertEqual(self.pipeline._detect_language("GraphRAG는 좋습니다"), "ko")

    def test_lancedb_table_creation(self):
        self.pipeline.initialize()
        table = self.pipeline._get_or_create_table()
        self.assertIsNotNone(table)

    def test_lancedb_insert_and_search(self):
        self.pipeline.initialize()
        table = self.pipeline._get_or_create_table()
        dim = self.pipeline.config["embed"]["dim"]

        # Insert
        vec = np.random.randn(dim).astype(np.float32).tolist()
        table.add([{"uid": 0, "source": "test", "text": "Hello world", "vector": vec}])

        # Search
        q_vec = np.random.randn(dim).astype(np.float32).tolist()
        results = table.search(q_vec).limit(1).to_list()
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]["text"], "Hello world")

    def test_nlp_standalone(self):
        self.pipeline.initialize()
        docs = self.pipeline.load_documents(["data/documents/sample.txt"])
        self.pipeline._run_standalone_nlp(docs)

        store_path = pathlib.Path(self.pipeline.config["ent"]["store_path"])
        self.assertTrue(store_path.exists())

        entities = []
        with open(store_path, encoding="utf-8") as f:
            for line in f:
                if line.strip():
                    entities.append(json.loads(line))
        self.assertGreater(len(entities), 0)

    def test_entity_context(self):
        self.pipeline.initialize()
        docs = self.pipeline.load_documents(["data/documents/sample.txt"])
        self.pipeline._run_standalone_nlp(docs)

        ctx = self.pipeline._get_entity_context("What is NLP?")
        # Should find something since NLP appears in the text
        self.assertIsInstance(ctx, str)


if __name__ == "__main__":
    unittest.main(verbosity=2)
